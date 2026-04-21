import streamlit as st
from docx import Document
import tempfile

# --- твојата мапа ---
MAC_C_TIMES_TO_UTF8 = {
    'a': 'а','b': 'б','v': 'в','g': 'г','d': 'д','|': 'ѓ','e': 'е','`': 'ж',
    'z': 'з','y': 'ѕ','i': 'и','j': 'ј','k': 'к','l': 'л','q': 'љ','m': 'м',
    'n': 'н','w': 'њ','o': 'о','p': 'п','r': 'р','s': 'с','t': 'т','}': 'ќ',
    'u': 'у','f': 'ф','h': 'х','c': 'ц','~': 'ч','x': 'џ','{': 'ш',

    'A': 'А','B': 'Б','V': 'В','G': 'Г','D': 'Д',']': 'Ќ','E': 'Е','@': 'Ж',
    'Z': 'З','Y': 'Ѕ','I': 'И','J': 'Ј','K': 'К','L': 'Л','Q': 'Љ','M': 'М',
    'N': 'Н','W': 'Њ','O': 'О','P': 'П','R': 'Р','S': 'С','T': 'Т','U': 'У',
    'F': 'Ф','H': 'Х','C': 'Ц','^': 'Ч','X': 'Џ','[': 'Ш',
}

def convert_text(text):
    return ''.join(MAC_C_TIMES_TO_UTF8.get(c, c) for c in text)

def process_docx(input_bytes):
    doc = Document(input_bytes)

    for paragraph in doc.paragraphs:
        if paragraph.text:
            converted = convert_text(paragraph.text)
            for run in paragraph.runs:
                run.text = ''
            if paragraph.runs:
                paragraph.runs[0].text = converted

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph.text:
                        converted = convert_text(paragraph.text)
                        for run in paragraph.runs:
                            run.text = ''
                        if paragraph.runs:
                            paragraph.runs[0].text = converted

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    doc.save(tmp.name)
    return tmp.name

# --- UI ---
st.title("📄 MAK Font Converter")

st.markdown("""

Оваа апликација конвертира Word (.docx) документи напишани со **Mac C Times** фонт 

во стандарден **UTF-8 македонски кириличен текст**.

### Како се користи:

1. Прикачете .docx документ

2. Почекајте обработка

3. Преземете го конвертираниот документ

ℹ️ Сите букви се мапираат автоматски според дефинирана табела.

""")
uploaded_file = st.file_uploader("Upload .docx file", type=["docx"])

if uploaded_file:
    with st.spinner("Processing..."):
        output_path = process_docx(uploaded_file)

    with open(output_path, "rb") as f:
        st.download_button(
            "⬇ Download converted file",
            f,
            file_name="converted.docx"
        )